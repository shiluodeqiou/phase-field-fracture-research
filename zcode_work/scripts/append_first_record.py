from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from pathlib import Path
p=Path(r'E:\909 研究\相场文献和程序\参考文献\相场文献阅读进展与问题记录.docx')
d=Document(p)
NAVY='17365D'; BLUE='D9EAF7'; PALE='F4F7FA'; BORDER='D9D9D9'
def runstyle(run,size=9.7,bold=False,color='000000'):
    run.font.name='Microsoft YaHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=RGBColor.from_string(color)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)
def border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); bs=tcPr.first_child_found_in('w:tcBorders')
    if bs is None: bs=OxmlElement('w:tcBorders'); tcPr.append(bs)
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=bs.find(qn('w:'+e))
        if x is None: x=OxmlElement('w:'+e); bs.append(x)
        x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),BORDER); x.set(qn('w:space'),'0')
def setc(c,text,head=False):
    c.text=text; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(c); shade(c,NAVY if head else PALE if False else 'FFFFFF')
    for par in c.paragraphs:
        par.paragraph_format.space_before=Pt(2); par.paragraph_format.space_after=Pt(2); par.paragraph_format.line_spacing=1.05
        for r in par.runs: runstyle(r,9.4,head,'FFFFFF' if head else '000000')

# Append a first-reading record.
h=d.add_paragraph(style='Heading 1'); h.paragraph_format.keep_with_next=True; r=h.add_run('第一篇阅读启动记录'); runstyle(r,15,True)
p0=d.add_paragraph(); p0.paragraph_format.space_after=Pt(6); r=p0.add_run('本次记录对应 2000 年的 Numerical experiments in revisited brittle fracture。当前完成的是整体导读和文章结构定位，不把它记为已经掌握全部推导。'); runstyle(r,10.5)

rows=[
 ('阅读日期','2026 年 10 月 27 日'),
 ('文献','B. Bourdin, G. A. Francfort, J.-J. Marigo, Numerical experiments in revisited brittle fracture'),
 ('阅读范围','摘要、引言、第 2 节数值实现的总体结构，以及第 3 节三个数值实验的结果定位'),
 ('文章主线','把 Francfort 和 Marigo 的变分脆性断裂模型转化为可计算的数值问题，并检验未知裂纹路径、裂纹起裂、渐进增长和分叉能否由能量框架自然给出'),
 ('已建立的理解','裂纹是未知集合；总能量由弹性能和裂纹表面能组成；裂纹随时间不可逆增长；在离散时间内，每一步都要在已有裂纹的基础上寻找当前能量更低的状态'),
 ('两种数值路线','一是引入辅助场 v 的椭圆正则化，v 接近 0 表示裂纹、接近 1 表示完整材料；二是使用自适应有限元，让网格本身贴合裂纹。两者都通过 Γ 收敛把尖锐裂纹问题转成可计算的近似问题'),
 ('尚待推导','为什么离散时间条件可以写成“裂纹不可逆加当前能量最小”；正则化能量中 c、k_c、h 的作用；交替最小化如何得到 u 和 v；Γ 收敛在这里保证了什么'),
 ('数值实验重点','加固体撕裂、纤维增强基体受拉、预裂纹分叉。文章重点观察裂纹形态和总能量变化是否符合理论预期，而不是完成与真实实验的定量标定'),
 ('下一步','精读引言中的能量和演化条件，再逐项推导第 2.1 节公式（1）至（3），最后回看第 3.1 节临界位移 δ_c 的计算')
]
t=d.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for i,x in enumerate(['记录项目','本次内容']): setc(t.rows[0].cells[i],x,True)
for a,b in rows:
    cells=t.add_row().cells; setc(cells[0],a); setc(cells[1],b)
for row in t.rows:
    row.cells[0].width=Pt(90); row.cells[1].width=Pt(420)

d.save(p)
print(p)
