from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

path = Path(r"E:/909 研究/相场文献和程序/参考文献/相场文献阅读进展与问题记录.docx")
doc = Document(path)

NAVY = '17365D'
BORDER = 'D9D9D9'
FONT = 'Microsoft YaHei'


def style_run(run, size=9.2, bold=False, color='000000'):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    bs = tcPr.first_child_found_in('w:tcBorders')
    if bs is None:
        bs = OxmlElement('w:tcBorders')
        tcPr.append(bs)
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = bs.find(qn('w:' + e))
        if x is None:
            x = OxmlElement('w:' + e)
            bs.append(x)
        x.set(qn('w:val'), 'single')
        x.set(qn('w:sz'), '6')
        x.set(qn('w:space'), '0')
        x.set(qn('w:color'), BORDER)


def set_cell(cell, text, size=9.0, bold=False, head=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)
    shade(cell, NAVY if head else 'FFFFFF')
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            style_run(r, size, True if head else bold, 'FFFFFF' if head else '000000')


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, 15 if level == 1 else 12, True, '000000')
    return p


def add_body(text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    style_run(r, size)
    return p


def add_bullet(text, size=10.5):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    style_run(r, size)
    return p


def add_table(headers, rows, widths, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, font_size, head=True)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell(cells[i], v, font_size)
        trPr2 = t.rows[-1]._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit')
        trPr2.append(cs)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    return t


# Update the overview status.
for row in doc.tables[0].rows[1:]:
    if '重访脆性断裂' in row.cells[1].text:
        set_cell(row.cells[4], '精读中：第 2 节（概念已通，推导待补）', 8.3)
        break

# Append the progress summary.
doc.add_page_break()
add_heading('阶段小结：第 1–2 节的进展与问题（2026-10-28）', 1)
add_body('本小结根据主对话和辅助对话的问答整理，区分“已读”“能复述”“待推导”“待核对”“待读”几个层次，避免把听过讲解记为已掌握。当前读的是 Bourdin–Francfort–Marigo (2000) 这篇文章，阅读顺序按发表时间从 2000 年开始。')

add_heading('一、阅读进展总表', 2)
add_table(['部分', '已读范围', '当前理解程度', '待补或备注'], [
    ('题目、摘要', 'PDF 正文 p.1', '能说明本文是 1998 年模型的数值实现，不是新模型；能说明 antiplane/plane 算例和“beyond classical tools”的含义。', '无。'),
    ('第 1 节', 'PDF 正文 pp.1–3', '能解释总能量、演化条件 (a)(b)(c)、时间离散、两条数值路线；能说明全局最小化是“最坏情况”。', 'inf 与 min 的数学细节；条件 (c) 的选择作用；离散条件 (b) 的符号差异（已讲解，暂不作为错误）。'),
    ('第 2 节', 'PDF 正文 pp.4–14', '能用“橡皮膜撕口 / 健康度 v / 模糊裂纹 / 两个编辑交替 / 渔网网格”讲清 AT 正则化和自适应网格的思想；知道 c、k_c、h 的作用。', 'SBV 与 Γ 收敛的证明；Euler–Lagrange 推导；一维剖面的积分；参数层级；自适应网格对偶化。'),
    ('第 3 节', 'PDF 正文 pp.14–21', '未读。', '3.1 撕裂；3.2 纤维增强基体；3.3 裂纹分叉。'),
], [0.95, 1.15, 2.55, 1.85], 8.5)

add_heading('二、目前能用自己的话讲清的内容', 2)
for text in [
    '本文不是提出新模型，而是把 1998 年 Francfort–Marigo 的变分断裂模型算出来。',
    '总能量 = 表面能 + 弹性能；裂纹路径由两者竞争决定，不需要额外的扩展准则。',
    '裂纹不可逆；每个载荷步在“包含上一步裂纹”的集合中最小化总能量。',
    '裂纹集合无法直接枚举；用 AT 正则化（连续场 v）或自适应网格（网格 T）表示。',
    'AT 正则化：v=1 完整、v=0 裂纹；c 控制模糊带宽度；k_c 防止刚度矩阵奇异；c→0 时收敛到尖锐裂纹。',
    '交替最小化：固定 v 解 u，固定 u 解 v；能量单调下降，但没有收敛证明，可能停在局部极小。',
    '方法 2 把网格本身当第二场，贴着裂纹走；更快，但能量估计和稳定性不如方法 1。',
    '平面弹性的非穿透条件仍未解决，所以平面算例只能定性看。',
]:
    add_bullet(text)

add_heading('三、优先补推导的问题', 2)
add_body('这五个问题是从“能讲清思想”走向“能独立推导”的关键。建议按顺序补，不要一次全啃。')
add_table(['编号', '具体问题', '位置', '类型', '状态', '目前的理解或待核对点'], [
    ('1', 'E_e 为什么是 inf 不是 min？', 'p.2 / p.6', '理论', '待推导', 'Γ 不够光滑时 inf 可能取不到；SBV 弱形式保证最小值存在。需要补 SBV 紧性。'),
    ('2', '条件 (c) 相对 (b) 增加了什么？', 'p.2', '理论', '待推导', 'Γ(s) 不包含 s 之后的裂纹，不在 (b) 的容许集里；它是“不后悔”条件。需要用一个具体例子验证。'),
    ('3', 'k_c << c 在 Γ 收敛中起什么作用？', 'p.7', '理论', '待推导', '需要理解上下极限的构造；k_c 不能破坏表面能极限。'),
    ('4', '交替最小化没有收敛证明；它收敛到什么？v0 如何影响？', 'pp.8–10', '数值', '待推导', '能量单调下降；子列收敛到临界点；非凸性导致局部极小。'),
    ('5', '一维剖面 v=1−exp(−|x|/2c) 为什么积分后表面能为 1？', 'p.9', '理论', '待推导', '每侧 1/2，两侧合计 1；1/(4c) 是归一化因子。'),
], [0.42, 1.90, 0.95, 0.75, 0.85, 1.63], 8.3)

add_heading('四、待核对或待理解的问题', 2)
add_body('这些问题不一定要马上推导，但需要在读第 3 节和回看第 2 节时逐条核对。')
add_table(['编号', '具体问题', '位置', '类型', '状态', '目前的理解或待核对点'], [
    ('6', '离散条件 (b) 原文 Γ⊃Γ_i，算法用 Γ⊃Γ_{i-1}，是否笔误？', 'p.3 / p.9', '文本核对', '已讲解，待核对', '原文确实印 Γ⊃Γ_i；它比直接离散化弱；算法用更强的条件；不影响主线。'),
    ('7', '全局最小化为什么是“最坏情况”？', 'p.3', '理论', '待推导', '全局最小化假设材料能跳出亚稳态；实际材料可能被能垒困住。'),
    ('8', 'v=1 完整 / v=0 裂纹，与现代 d=0 完整 / d=1 裂纹如何对应？', 'pp.3–4 / §2.1', '符号', '已讲解，待复述', 'd=1−v。'),
    ('9', '理论 h << k_c << c 与算例 h=1e-3、k_c=1e-4 的关系？', 'pp.8–9', '数值', '待核对', '算例满足 h<<c 和 k_c<<c，但 h/k_c=10，不在严格渐近层级；实际裂纹附近只需 h<<c。'),
    ('10', 'Ω̃ 要取多大？误差为什么是 O(exp(−T/c))？', 'pp.9–10', '数值', '待推导', 'T=dist(S_u,∂Ω̃)；需要 T 是 c 的若干倍。'),
    ('11', '自适应网格的对偶化：网格 T 和 v_T 各自起什么作用？', 'pp.11–13', '数值', '待推导', '网格是第二场；v_T 是单元常数；M_T(v) 由 f′ 显式给出。'),
    ('12', '平面弹性为什么不能直接沿用反平面结论？非穿透为什么是开放问题？', 'p.14', '理论', '待读/待推导', '强/弱形式等价未建立；裂纹面互不穿透需要单边约束；图 4(f) 出现非物理结果。'),
], [0.42, 1.95, 0.95, 0.75, 0.85, 1.58], 8.3)

add_heading('五、下一步阅读', 2)
add_table(['编号', '待读内容', '位置', '类型', '状态', '阅读时要核对什么'], [
    ('13', '第 3 节三个算例', 'PDF pp.14–21', '算例', '待读', '3.1 临界位移 δ_c=√(2 log 3)≈1.48、表面能 2π、能量曲线；3.2 Griffith 准则、对称性破缺；3.3 分叉、非穿透。'),
], [0.42, 1.55, 1.05, 0.65, 0.55, 2.28], 8.4)

add_heading('六、学习方式备注', 2)
for text in [
    '第 2 节原文密度高，先抓思想，再抠公式；当前最有效的比喻是“橡皮膜撕口 + 健康度 v + 模糊裂纹 + 两个编辑交替 + 渔网网格”。',
    '目前对第 2 节的状态是“能讲清思想，但还不能独立推导”；不要把它记为“已掌握”。',
    '下一步先读 3.1，用具体算例检验抽象模型；读完后再回到第 2 节补 SBV、Γ 收敛和参数层级的推导。',
    '每次更新保留“已讲解”和“已掌握”的区别；问题解决后保留原问题文字，再补充依据。',
]:
    add_bullet(text)

doc.save(path)
print('saved:', path)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
