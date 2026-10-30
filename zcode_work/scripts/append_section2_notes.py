from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

path = Path(r"E:/909 研究/相场文献和程序/参考文献/相场文献阅读进展与问题记录.docx")
doc = Document(path)

NAVY = '17365D'
BORDER = 'D9D9D9'
FONT = 'Microsoft YaHei'


def style_run(run, size=9.2, bold=False, color='000000'):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    bs = tcPr.first_child_found_in('w:tcBorders')
    if bs is None:
        bs = OxmlElement('w:tcBorders')
        tcPr.append(bs)
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = bs.find(qn('w:' + e))
        if x is None:
            x = OxmlElement('w:' + e)
            bs.append(x)
        x.set(qn('w:val'), 'single')
        x.set(qn('w:sz'), '6')
        x.set(qn('w:space'), '0')
        x.set(qn('w:color'), BORDER)


def set_cell(cell, text, size=9.0, bold=False, head=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)
    shade(cell, NAVY if head else 'FFFFFF')
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            style_run(r, size, True if head else bold, 'FFFFFF' if head else '000000')


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, 15 if level == 1 else 12, True, '000000')
    return p


def add_body(text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    style_run(r, size)
    return p


def add_table(headers, rows, widths, font_size=8.6):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, font_size, head=True)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell(cells[i], v, font_size)
        trPr2 = t.rows[-1]._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit')
        trPr2.append(cs)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    return t


# Update the status of the first paper.
for row in doc.tables[0].rows[1:]:
    if '重访脆性断裂' in row.cells[1].text:
        set_cell(row.cells[4], '精读中：第 2 节', 8.6)
        break

# Append the new reading note.
doc.add_page_break()
add_heading('第一篇精读笔记（二）：第 2 节 数值实现', 1)
add_body('本次记录对应第 2 节（PDF 正文 pp.4–14；期刊 pp.800–810）。第 2 节把尖锐裂纹的自由不连续问题转化为可计算的 AT 正则化或自适应网格问题，并用 Γ 收敛说明近似模型与原来模型的极限关系。本次只记录主线、关键公式和仍未解决的问题，不把每一步推导都记为已掌握。')

add_heading('一、本次阅读卡片', 2)
add_table(['项目', '内容'], [
    ('阅读日期', '2026 年 10 月 28 日'),
    ('文献', 'B. Bourdin, G. A. Francfort, J.-J. Marigo, Numerical experiments in revisited brittle fracture, J. Mech. Phys. Solids 48(4) (2000) 797–826.'),
    ('阅读范围', '第 2 节（PDF 正文 pp.4–14；期刊 pp.800–810）'),
    ('本次目标', '搞清强/弱形式、SBV、AT 正则化、Γ 收敛、交替最小化、参数选择、自适应有限元'),
    ('一句话概括', '尖锐裂纹的自由不连续问题无法直接有限元离散；用 AT 正则化或自适应网格把它变成可计算的场/网格问题，并用 Γ 收敛保证极限一致。'),
    ('已理解', '强形式与弱形式的关系；SBV 的作用；E_c 的两个积分；一维指数剖面；交替最小化流程；k_c、c、h 的作用；两种方法的优缺点。'),
    ('尚未理解', 'k_c << c 在 Γ 收敛证明中的具体作用；理论层级 h << k_c << c 与实际参数 h=1e-3、k_c=1e-4 的关系；交替最小化收敛到什么临界点；Ω̃ 大小的误差估计；自适应网格的对偶化细节。'),
    ('下一步', '第 3 节三个算例，先读 3.1 撕裂问题，重点核对临界位移 δ_c、能量曲线和表面能误差。'),
], [1.35, 5.15], 9.0)

add_heading('二、关键公式', 2)
add_table(['编号', '公式', '说明'], [
    ('(1)', 'F(u,v) = E(u)（u∈SBV(Ω̃)，u=U on Ω̃∖Ω̄，v=1 a.e.）；否则 +∞', '把单场能量扩展成两场泛函；+∞ 用来强制 v=1。'),
    ('(2)', 'E_c(u,v) = ∫_{Ω∖N}(v²+k_c)|∇u|² dx + ∫_{Ω̃∖N}(c|∇v|² + (1−v)²/(4c)) dx', '第一项是退化弹性能；第二项是表面能正则化。'),
    ('(3)', 'F_c(u,v) = E_c（u∈W^{1,2}(Ω̃∖N)，u=U on Ω̃∖Ω̄，v∈W^{1,2}(Ω̃∖N;[0,1])）；否则 +∞', '定义正则化泛函的容许集。'),
    ('离散', 'E_{c,h}：E_c 在分片仿射有限元空间上的投影', '若 h 比 k_c、c 更快趋于 0，则同样的 Γ 收敛成立；理论层级 h << k_c << c。'),
    ('(9)', '∑_{T∈T_h}|T∩Ω|(v_T M_T^*(|∇u|²) + ψ(v_T)/h_T)', '自适应网格法的对偶泛函；v_T 是单元常数。'),
    ('(10)', '∫_{Ω∖N}(v²+k_c)W(e(u))dx + ∫_{Ω̃∖N}(c|∇v|² + (1−v)²/(4c))dx', '平面弹性的正则化泛函；W(ξ)=1/2{tr(ξ)²+2ξ:ξ}。'),
], [0.55, 2.65, 3.30], 8.4)

add_heading('三、交替最小化算法（§2.1.1）', 2)
add_table(['步骤', '内容'], [
    ('初始化', '固定 c、h、k_c；在 Ω̃ 上建网格；选初始 v0。文中在每个载荷步取 v0 ≡ 1。'),
    ('迭代 k 第 1 步', '固定 v_{k-1}，在 Ω 上最小化 E_{c,h}(·,v_{k-1})，得到 u_k。'),
    ('迭代 k 第 2 步', '把 u_k 按边界值 U 延拓到 Ω̃∖Ω。'),
    ('迭代 k 第 3 步', '固定 u_k，在 Ω̃ 上最小化 E_{c,h}(u_k,·)，得到 v_k。'),
    ('迭代 k 第 4 步', '若 ||v_k − v_{k-1}||∞ ≥ ε，继续；否则停止。'),
    ('下一个载荷步', '在已检测到裂纹的节点上强制 v=0（不可逆），再重新初始化或沿用上一步的 v。'),
    ('计算代价', '单个载荷步可能需要进行多达 150 次交替迭代。'),
    ('收敛性', '能量单调下降；没有收敛证明；子列收敛到 E_{c,h} 的某个临界点。'),
], [1.45, 5.05], 8.8)

add_heading('四、符号与概念对照', 2)
add_table(['概念', '含义', '易错点'], [
    ('SBV', '特殊有界变差函数空间；允许跳跃；跳跃集 S_u 就是裂纹。', '普通 W^{1,2} 不能表示位移跳跃，所以必须换空间。'),
    ('Γ 收敛', '泛函序列的变分收敛；极小元收敛到极限泛函的极小元。', '只保证极限，不给误差估计。'),
    ('AT 正则化', '用连续场 v 代替尖锐裂纹；v=1 完整，v=0 裂纹。', '本文 v 是完整度；现代相场常用 d=0 完整、d=1 裂纹，d=1−v。'),
    ('逻辑域 Ω̃', '包含 Ω̄ 的更大区域；位移在外边界给定。', '允许裂纹到达物理边界；表面能仍算在 Ω̄∖N。'),
    ('参数 k_c', '残余刚度，防止 v=0 时刚度矩阵奇异。', '太大会保留残余刚度，高估弹性能；太小会数值发散。'),
    ('参数 c', '裂纹过渡带宽度（正则化长度）。', '太大有软化、低估弹性能；太小相对 h 会高估表面能。'),
    ('网格 h', '有限元尺寸。', '裂纹附近需要 h << c；理论上 h 比 k_c、c 更快趋于 0。'),
    ('方法 2', '第二个场是网格 T；G_h 直接定义在三角剖分上。', '对偶变量 v_T 是单元常数；网格自适应贴着裂纹走。'),
], [1.10, 2.45, 2.95], 8.4)

add_heading('五、本节新增问题', 2)
add_table(['编号', '具体问题', '关联位置', '问题类型', '状态', '目前的理解或待核对点'], [
    ('1', '为什么弱形式要用 SBV，而不能继续用 W^{1,2}？', 'p.5', '理论', '待解决', 'W^{1,2} 不能表示跨裂纹的位移跳跃；SBV 的跳跃集可以表示裂纹。'),
    ('2', 'Γ 收敛为什么要求 k_c << c？k_c 在证明中起什么作用？', 'p.7', '理论', '待解决', '需要理解下半极限和上半极限的构造；k_c 不能破坏表面能极限。'),
    ('3', '文中说 h 比 k_c、c 更快趋于 0，但算例取 h=1e-3、k_c=1e-4。算例是否不在渐近层级内？', 'pp.8–9', '理论/数值', '待核对', '理论层级是 h << k_c << c；算例只满足 h << c 和 k_c << c，h/k_c=10。'),
    ('4', '交替最小化没有收敛证明；它实际收敛到什么？初始 v0 如何影响结果？', 'pp.8–10', '数值', '待解决', '能量单调下降；子列收敛到临界点；非凸性导致局部极小。'),
    ('5', '一维剖面 v=1−exp(−|x|/2c) 为什么积分后表面能是 1？', 'p.9', '理论', '待解决', '每侧积分 1/2，两侧合计 1；因子 1/(4c) 的选择。'),
    ('6', '逻辑域 Ω̃ 要取多大？误差为什么是 O(exp(−T/c))？', 'pp.9–10', '数值', '待解决', 'T=dist(S_u,∂Ω̃)；需要 T 是 c 的若干倍。'),
    ('7', '自适应网格法中，网格 T 和对偶变量 v_T 各自起什么作用？', 'pp.11–13', '数值', '待解决', '网格是第二场；v_T 是单元常数；M_T(v) 由 f′ 显式给出。'),
    ('8', '平面弹性为什么不能直接沿用反平面结论？非穿透条件为什么是开放问题？', 'p.14', '理论', '待解决', '强/弱形式等价未建立；裂纹面互不穿透需要单边约束；图 4(f) 出现非物理结果。'),
], [0.42, 2.00, 1.05, 0.85, 0.55, 1.63], 8.2)

add_heading('六、本次形成的五句话', 2)
for text in [
    '第 2 节的核心不是新模型，而是把尖锐裂纹的自由不连续问题变成可计算的近似问题。',
    'SBV 弱形式是 AT 正则化和 Γ 收敛的出发点；普通有限元无法直接处理未知跳跃集。',
    'AT 正则化用 v 表示完整度：v=1 完整、v=0 裂纹；c 控制过渡带，k_c 防止刚度奇异。',
    '交替最小化是局部下降，没有收敛证明；能量单调下降，但可能停在局部极小。',
    '自适应网格法把网格本身当作第二场；它更快，但能量估计和稳定性不如 AT 方法。',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    style_run(r, 10.5)

add_heading('七、下一步', 2)
add_body('下一次进入第 3 节，先读 3.1 加固体撕裂：核对临界位移 δ_c = √(2 log 3) ≈ 1.48、表面能 2π、能量曲线，以及计算临界载荷偏低的原因。之后再读 3.2 纤维增强基体和 3.3 裂纹分叉。')

doc.save(path)
print('saved:', path)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
