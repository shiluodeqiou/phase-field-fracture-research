from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

out = Path(r"E:\909 研究\相场文献和程序\参考文献\相场文献阅读进展与问题记录.docx")

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

# Color palette
NAVY = '17365D'
BLUE = 'D9EAF7'
PALE = 'F4F7FA'
BORDER = 'D9D9D9'
BLACK = '000000'
GRAY = '666666'

# helpers
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=BORDER, sz='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), sz)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_col_widths(table, widths):
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)

def format_table(table, widths, header=True, font_size=9.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_col_widths(table, widths)
    for ri, row in enumerate(table.rows):
        if header and ri == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and ri == 0:
                set_cell_shading(cell, NAVY)
            elif ri % 2 == 0:
                set_cell_shading(cell, PALE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = 'Microsoft YaHei'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                    run.font.size = Pt(font_size)
                    if header and ri == 0:
                        run.font.color.rgb = RGBColor(255,255,255)
                        run.bold = True
                    else:
                        run.font.color.rgb = RGBColor(0,0,0)

def add_heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Microsoft YaHei'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    r.font.color.rgb = RGBColor(0,0,0)
    return p

def add_body(text, bold_prefix=None):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = 'Microsoft YaHei'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0,0,0)
    return p

def add_field_table(rows, widths=(1.45,5.05), font_size=10):
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    format_table(table, widths, header=False, font_size=font_size)
    for row in table.rows:
        set_cell_shading(row.cells[0], BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
    return table

# Styles
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(10.5)
for name, size, color in [('Title', 22, BLACK), ('Heading 1', 15, BLACK), ('Heading 2', 12, BLACK)]:
    st = styles[name]
    st.font.name = 'Microsoft YaHei'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(12 if name!='Title' else 0)
    st.paragraph_format.space_after = Pt(6)

# Footer
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('相场文献阅读记录  ·  可持续更新')
run.font.name = 'Microsoft YaHei'
run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(102,102,102)

# Title block
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_after = Pt(4)
r = p.add_run('相场文献阅读进展与问题记录')
r.font.name = 'Microsoft YaHei'
r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
r.font.color.rgb = RGBColor(0,0,0)
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(18)
r = sub.add_run('研究主题：断裂相场模型、氢致开裂、机器学习与程序实现')
r.font.name = 'Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); r.font.size = Pt(11); r.font.color.rgb = RGBColor(89,89,89)

add_body('本记录用于持续整理相场断裂文献的阅读进展、关键理解、尚未解决的问题以及与现有程序的对应关系。当前目录已经整理出 7 篇核心文献，并保留了 PDF 与 Markdown 文本；后续每次阅读只需在相应记录中补充内容，并在问题清单中追踪问题的解决状态。')

add_heading('使用方法', 1)
for t in [
    '每次阅读先填写日期、文献和阅读范围，再用自己的话概括模型假设、能量泛函、控制方程和数值方法。',
    '公式不必全部抄录，但要记录每个变量的物理意义、边界条件、不可逆条件以及与程序输入或子程序的对应位置。',
    '遇到不理解的内容，先写出具体问题和所在页码或公式编号；问题解决后补充依据、推导或程序验证结果。',
    '建议按照“基础理论 → 数值实现 → 专题扩展 → 程序复现”的顺序推进，不把只看过摘要或听过讲解记为已掌握。'
]:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    p.add_run(t)
    for run in p.runs:
        run.font.name='Microsoft YaHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); run.font.size=Pt(10.5)

add_heading('文献总览', 1)
add_body('下表根据当前“参考文献”目录中的文件整理。建议阶段是阅读顺序建议，当前状态留作后续更新。')
headers = ['序号','文献','研究主题','建议阶段','当前状态']
rows = [
    ['1','重访脆性断裂中的数值实验\nNumerical experiments in revisited brittle fracture','变分断裂与数值实验','基础','待开始'],
    ['2','热力学一致的断裂相场模型\nThermodynamically consistent phase-field models of fracture','热力学、变分原理、多场有限元','基础','待开始'],
    ['3','断裂的相场建模\nPhase-field modelling of fracture','模型综述、理论与实现','基础','待开始'],
    ['4','脆性断裂相场模型综述与一种新的快速混合列式\nA review on phase-field models of brittle fracture and a new fast hybrid formulation','综述、混合求解与效率','数值','待开始'],
    ['5','韧性断裂的相场建模\nPhase-field modeling of ductile fracture','塑性与韧性断裂','专题','待开始'],
    ['6','氢致开裂的相场列式\nA phase field formulation for hydrogen assisted cracking','氢扩散、力学与断裂耦合','专题','待开始'],
    ['7','结合物理信息深度学习的断裂相场建模\nPhase-field modeling of fracture with physics-informed deep learning','深度 Ritz 方法与机器学习','专题','待开始'],
]
t = doc.add_table(rows=1, cols=len(headers))
for i,h in enumerate(headers): t.rows[0].cells[i].text=h
for row in rows:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v
format_table(t, [0.38,2.45,1.45,0.78,0.80], header=True, font_size=8.7)

add_heading('当前起点记录', 1)
add_body('这条记录只描述当前资料整理阶段，后续应以实际阅读和程序验证结果为准。')
add_field_table([
    ('记录日期','2026 年 10 月 27 日'),
    ('本次范围','完成文献目录整理；确认 7 篇文献均有 PDF，并有对应的 Markdown 文本'),
    ('已完成内容','在 Zotero 中完成标题汉化并导入“相场文献和程序”栏目；本地目录中已按“参考文献”和“相场程序”分开保存'),
    ('当前进展','已具备统一记录入口；尚未将“读过”视为“掌握”，需要逐篇完成公式、假设和程序对应关系的核对'),
    ('当前问题','各文献采用的相场变量约定、裂纹不可逆处理、张拉压缩分裂方式和求解算法之间有什么差异？现有程序分别对应哪一种模型？'),
    ('下一步','先阅读综述和热力学一致模型，建立统一符号表；随后选一篇程序相关文献做最小算例复现')
], widths=(1.45,5.05), font_size=9.7)

add_heading('单篇文献记录模板', 1)
add_body('每完成一次阅读，可复制下面的表格作为新的记录。页码、公式编号和程序文件名尽量写具体，便于之后回查。')
add_field_table([
    ('阅读日期',''),
    ('文献及版本',''),
    ('阅读范围','章节、页码或公式编号'),
    ('阅读目标','本次希望回答什么问题'),
    ('核心问题','作者要解决什么问题，采用了什么模型'),
    ('模型假设','材料、几何、载荷、裂纹和场变量的主要假设'),
    ('关键公式','能量泛函、控制方程、边界条件、不可逆条件及变量含义'),
    ('数值方法','离散方式、求解流程、收敛准则、网格或长度尺度要求'),
    ('我的理解','用自己的话写出本次真正理解的内容'),
    ('尚未理解','具体到页码、公式、图或术语的问题'),
    ('程序对应','对应的目录、文件、子程序、输入参数或待验证位置'),
    ('验证结果','推导、手算、算例、代码运行或图表核对结果'),
    ('下一步','下一次阅读前要完成的一个或多个小任务')
], widths=(1.45,5.05), font_size=9.6)

add_heading('问题清单', 1)
add_body('问题写成可检查的句子，例如“公式（35）中的历史变量是否等价于程序中的 HMAX？”；解决后填写依据，不只写“已解决”。')
headers = ['编号','具体问题','关联文献或位置','问题类型','状态','解决依据或验证结果']
t = doc.add_table(rows=1, cols=len(headers))
for i,h in enumerate(headers): t.rows[0].cells[i].text=h
for i in range(1,9):
    cells=t.add_row().cells
    vals=[str(i),'','','理论 / 数值 / 程序','','']
    for j,v in enumerate(vals): cells[j].text=v
format_table(t, [0.38,2.20,1.30,0.90,0.55,1.17], header=True, font_size=8.6)

add_heading('阶段计划', 1)
headers = ['阶段','目标','完成标志','计划时间','实际记录']
t = doc.add_table(rows=1, cols=len(headers))
for i,h in enumerate(headers): t.rows[0].cells[i].text=h
plan = [
    ['第一阶段','读懂综述和变分基础','能解释相场变量、裂纹表面正则化、能量分解和不可逆条件','待填写',''],
    ['第二阶段','读懂有限元实现','能画出位移场与相场的求解流程，并说明每个残量的来源','待填写',''],
    ['第三阶段','专题模型对比','完成韧性断裂、氢致开裂和机器学习模型的差异表','待填写',''],
    ['第四阶段','程序复现与诊断','至少完成一个最小算例，并记录输入、输出、误差和未解决问题','待填写',''],
]
for row in plan:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v
format_table(t, [0.82,1.65,2.55,0.72,0.76], header=True, font_size=8.7)

add_heading('文献与程序对应关系', 1)
add_body('以下对应关系依据当前目录名称作初步整理，首次运行或精读相关文献时应再次核对。')
headers=['程序目录或文件','可能对应的研究内容','需要核对的文献','核对结果']
t=doc.add_table(rows=1, cols=4)
for i,h in enumerate(headers): t.rows[0].cells[i].text=h
mapping=[
    ['程序目录1 相场模型 氢致裂纹','力学 氢扩散 相场断裂耦合；含 PhaseField.for、Job-1.inp、ResultsQ4.py','氢致开裂的相场列式','待验证'],
    ['程序目录2 相场模型 机器学习','物理信息深度学习相场断裂；含 phase-field-fracture-with-pidl-main.zip','结合物理信息深度学习的断裂相场建模','待验证'],
    ['程序目录3 相场模型','通用相场断裂程序；含 GRIPHFiTH-devel.zip','综述、快速混合列式或相关实现','待验证'],
]
for row in mapping:
    cells=t.add_row().cells
    for i,v in enumerate(row): cells[i].text=v
format_table(t, [1.55,2.30,1.85,0.80], header=True, font_size=8.8)

add_heading('更新规则', 1)
add_body('每次更新在记录日期后新增一条单篇文献记录；问题解决后保留原问题文字，再补充解决依据。这样可以区分“读过”“理解”“已验证”三个层次，也方便向老师汇报实际进展。')

# document core properties
doc.core_properties.title = '相场文献阅读进展与问题记录'
doc.core_properties.subject = '断裂相场模型文献阅读与程序复现记录'
doc.core_properties.author = '学生'
doc.core_properties.comments = '根据当前参考文献和相场程序目录建立的可持续更新模板'

doc.save(out)
print(out)
