from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

path = Path(r"E:/909 研究/相场文献和程序/参考文献/相场文献阅读进展与问题记录.docx")
doc = Document(path)

# Prevent every table row from splitting across pages.
for t in doc.tables:
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            cantSplit = OxmlElement('w:cantSplit')
            tblHeader = trPr.find(qn('w:tblHeader'))
            if tblHeader is not None:
                tblHeader.addprevious(cantSplit)
            else:
                trPr.append(cantSplit)

# Widen the overview table to the full text width and give the number column room.
overview_widths = [0.45, 2.65, 1.55, 0.85, 1.00]
for row in doc.tables[0].rows:
    for i, w in enumerate(overview_widths):
        row.cells[i].width = Inches(w)
doc.tables[0].autofit = False

# Give the problem-list number column a little more room.
problem_widths = [0.48, 2.10, 1.30, 0.90, 0.55, 1.17]
for row in doc.tables[3].rows:
    for i, w in enumerate(problem_widths):
        row.cells[i].width = Inches(w)
doc.tables[3].autofit = False

doc.save(path)
print('saved:', path)
