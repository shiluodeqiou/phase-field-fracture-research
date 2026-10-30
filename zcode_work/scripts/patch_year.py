from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from pathlib import Path
p=Path(r'E:\909 研究\相场文献和程序\参考文献\相场文献阅读进展与问题记录.docx')
d=Document(p)
cell=d.tables[0].rows[4].cells[3]
cell.text='2018（图示标注）\nMarkdown 标注 2019'
for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(2); par.paragraph_format.space_before=Pt(2); par.paragraph_format.line_spacing=1.05
    for run in par.runs:
        run.font.name='Microsoft YaHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); run.font.size=Pt(8.6)
d.save(p)
print(p)
