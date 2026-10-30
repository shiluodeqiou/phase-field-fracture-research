from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path

path = Path(r'E:\909 研究\相场文献和程序\参考文献\相场文献阅读进展与问题记录.docx')
img = Path(r'E:\909 研究\相场文献和程序\参考文献\主要参考文献说明.png')
doc = Document(path)

def style_run(run, size=9.5, bold=False, color='000000'):
    run.font.name='Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    run.font.size=Pt(size)
    run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)

def set_cell(cell, text, size=8.7, bold=False):
    cell.text = text
    for p in cell.paragraphs:
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.space_before=Pt(2)
        p.paragraph_format.line_spacing=1.05
        for run in p.runs: style_run(run,size,bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

# Replace the literature overview with the order and annotations shown in the supplied image.
t = doc.tables[0]
rows = [
 ['1','氢致开裂的相场列式\nA phase field formulation for hydrogen assisted cracking','相场模型与氢致裂纹','2018\n程序目录 1','待开始'],
 ['2','脆性断裂相场模型综述与一种新的快速混合列式\nA review on phase-field models of brittle fracture and a new fast hybrid formulation','脆性断裂相场方法与综述','2014（图示标注 2015）','待开始'],
 ['3','重访脆性断裂中的数值实验\nNumerical experiments in revisited brittle fracture','早期文献、偏理论推导与数值实验','2000','待开始'],
 ['4','断裂的相场建模\nPhase-field modelling of fracture','较详细的断裂相场综述','2018','待开始'],
 ['5','韧性断裂的相场建模\nPhase-field modeling of ductile fracture','韧性断裂相场方法与综述','2015','待开始'],
 ['6','结合物理信息深度学习的断裂相场建模\nPhase-field modeling of fracture with physics-informed deep learning','相场模型与机器学习','2024\n程序目录 2、3','待开始'],
 ['7','热力学一致的断裂相场模型\nThermodynamically consistent phase-field models of fracture: Variational principles and multi-field FE implementations','奠定计算基础的代表性工作','2010','待开始'],
]
for i,row in enumerate(rows, start=1):
    for j,val in enumerate(row):
        set_cell(t.rows[i].cells[j], val, size=8.6)
# Ensure header remains formatted.
for cell in t.rows[0].cells:
    set_cell(cell, cell.text, size=8.7, bold=True)
    for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255)

# Clarify the source note in the overview paragraph.
for p in doc.paragraphs:
    if p.text.startswith('下表根据当前'):
        p.text = '下表根据当前“参考文献”目录和你提供的“主要参考文献说明”图片整理。图片中的年份和代表性说明作为阅读顺序参考；若与 PDF 正文或出版信息不一致，记录中保留差异并以原始文献为准。建议阶段是阅读顺序建议，当前状态留作后续更新。'
        for run in p.runs: style_run(run,10.5)
        break

# Update the current starting record to include the supplied image as a source.
for table in doc.tables:
    for row in table.rows:
        if len(row.cells)>=2 and row.cells[0].text.strip()=='本次范围':
            set_cell(row.cells[1], '完成文献目录整理；确认 7 篇文献均有 PDF，并有对应的 Markdown 文本；补充核对“主要参考文献说明”图片中的年份、作者和程序目录标注', size=9.7)
        if len(row.cells)>=2 and row.cells[0].text.strip()=='当前问题':
            set_cell(row.cells[1], '各文献采用的相场变量约定、裂纹不可逆处理、张拉压缩分裂方式和求解算法之间有什么差异？现有程序分别对应哪一种模型？图片中综述文献的年份与 PDF 信息为何存在差异？', size=9.7)

# Update program mapping table (last table) to match the image.
tm = doc.tables[-1]
for i,row in enumerate([
 ['程序目录 1 相场模型 氢致裂纹','力学、氢扩散与相场断裂耦合；含 PhaseField.for、Job-1.inp、ResultsQ4.py','氢致开裂的相场列式','待验证'],
 ['程序目录 2、3 相场模型与机器学习','物理信息深度学习相场断裂；图片将其同时标注为程序目录 2 和 3','结合物理信息深度学习的断裂相场建模','待验证'],
 ['程序目录 3 相场模型','通用相场断裂程序；含 GRIPHFiTH-devel.zip','与综述或快速混合列式的具体对应关系待核对','待验证'],
], start=1):
    for j,val in enumerate(row): set_cell(tm.rows[i].cells[j], val, size=8.6)

# Append the supplied image and a short source note for auditability.
p = doc.add_paragraph()
p.paragraph_format.space_before=Pt(14)
p.paragraph_format.keep_with_next=True
r=p.add_run('主要参考文献说明图')
style_run(r,15,True)
p2=doc.add_paragraph()
p2.paragraph_format.space_after=Pt(6)
r=p2.add_run('该图作为文献优先级、代表性和程序目录标注的参考来源保存在“参考文献”目录中。图中信息用于组织学习路线，具体年份和出版信息继续以 PDF 或 DOI 页面为准。')
style_run(r,10.5)
p3=doc.add_paragraph()
p3.alignment=WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_after=Pt(4)
p3.add_run().add_picture(str(img), width=Inches(6.5))
cap=doc.add_paragraph()
cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=cap.add_run('图 1 主要参考文献说明（用户提供）')
style_run(r,9,False,'666666')

doc.save(path)
print(path)
