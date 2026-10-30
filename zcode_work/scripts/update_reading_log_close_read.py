from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

path = Path(r"E:/909 研究/相场文献和程序/参考文献/相场文献阅读进展与问题记录.docx")
doc = Document(path)

NAVY = '17365D'
BORDER = 'D9D9D9'
FONT = 'Microsoft YaHei'


def style_run(run, size=9.2, bold=False, color='000000'):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    bs = tcPr.first_child_found_in('w:tcBorders')
    if bs is None:
        bs = OxmlElement('w:tcBorders')
        tcPr.append(bs)
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = bs.find(qn('w:' + e))
        if x is None:
            x = OxmlElement('w:' + e)
            bs.append(x)
        x.set(qn('w:val'), 'single')
        x.set(qn('w:sz'), '6')
        x.set(qn('w:space'), '0')
        x.set(qn('w:color'), BORDER)


def set_cell(cell, text, size=9.0, bold=False, head=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)
    shade(cell, NAVY if head else 'FFFFFF')
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            style_run(r, size, True if head else bold, 'FFFFFF' if head else '000000')


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, 15 if level == 1 else 12, True, '000000')
    return p


def add_body(text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    style_run(r, size)
    return p


def add_table(headers, rows, widths, font_size=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, font_size, head=True)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell(cells[i], v, font_size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    return t


# 1) Update paper status in overview
overview = doc.tables[0]
for row in overview.rows[1:]:
    if '重访脆性断裂' in row.cells[1].text:
        set_cell(row.cells[4], '精读中：第 1 节', 8.6)
        break

# 2) Fill problem list (table index 3)
problems = doc.tables[3]
problem_rows = [
    ('1', '为什么弹性能定义为 inf 而不是 min？', '正文 p.2（期刊 p.798），E_e 定义', '理论', '待解决',
     '原文：Γ 不够光滑时 inf 可能取不到；弱形式用 SBV 保证最小值存在（p.6）。需要进一步理解 SBV 与紧性。'),
    ('2', '演化条件 (c) 相对 (b) 增加了什么？为什么它是选择准则？', '正文 p.2，条件 (a)-(c)', '理论', '待解决',
     'Γ(s) 不包含 s 之后的裂纹，因此不在 (b) 的容许集合中；它是防止“过早/过晚”扩展的选择条件，需要用具体例子验证。'),
    ('3', '离散条件 (b) 原文写 Γ ⊃ Γ_i，按连续条件离散化是否应为 Γ ⊃ Γ_{i-1}？', '正文 p.3；算法 p.9', '理论/文本核对', '待核对',
     '算法实际是在 Γ ⊃ Γ_{i-1} 上最小化；需查 Francfort-Marigo (1998) 原文确认印刷符号。'),
    ('4', '为什么全局最小化被称为“最坏情况”？局部极小值与亚稳态如何影响数值结果？', '正文 p.3', '理论/数值', '待解决',
     '全局最小化假设材料总能找到最低能量状态；实际材料可能被能垒困在亚稳态。非凸性也会使数值格式落在局部极小值。'),
    ('5', '本文 v=1 完整、v=0 裂纹，与现代相场 d=0 完整、d=1 裂纹如何对应？', '正文 p.3-4；§2.1', '符号/理论', '待解决',
     '初步：令 d=1-v；需核对 §2.1 的能量、刚度退化和边界条件。本文称 v 为 damage variable，但 v=1 是完整材料，符号容易混淆。'),
    ('6', 'Γ 收敛只保证极限，不给出误差估计；c、k_c、h 应如何选择？', '正文 p.8-9；§2.1.2', '数值', '待解决',
     '理论上 h 比 k_c、c 更快趋于 0，且 k_c << c；实际算例 h=10^-3、c=10^-1、k_c=10^-4。裂纹附近需要 h << c，k_c 主要防止刚度奇异。'),
    ('7', r'表面能为什么写成 k H¹(Γ\N)，为什么要减去 N？', '正文 p.2', '理论', '待解决',
     'H¹ 在二维中是裂纹长度；N 是 traction-free 边界，裂纹落在这部分不增加新的自由表面。'),
    ('8', '本文的 Γ 收敛与后来相场模型的 AT 正则化是什么关系？', '§2.1；p.4-6', '理论/文献', '待解决',
     '本文用 Ambrosio-Tortorelli 泛函；需要与后续相场文献中的 AT 模型逐项对照，尤其是 v 与 d 的符号和 1/(4c) 系数。'),
]
for i, vals in enumerate(problem_rows, start=1):
    if i >= len(problems.rows):
        break
    for j, v in enumerate(vals):
        set_cell(problems.rows[i].cells[j], v, 8.5)

# 3) Append close-reading section
doc.add_page_break()
add_heading('第一篇精读笔记（一）：题目、摘要与第 1 节', 1)
add_body('本次记录对应正文 pp.1–3（期刊 pp.797–799）。阅读方式是逐段核对原文：先看作者要解决什么问题，再看能量泛函和演化条件，最后看这些条件如何变成数值算法。第 2 节的公式推导尚未展开，因此状态记为“精读中”，不记为“已掌握”。')

add_heading('一、本次阅读卡片', 2)
add_table(['项目', '内容'], [
    ('阅读日期', '2026 年 10 月 28 日'),
    ('文献', 'B. Bourdin, G. A. Francfort, J.-J. Marigo, Numerical experiments in revisited brittle fracture, J. Mech. Phys. Solids 48(4) (2000) 797–826. DOI: 10.1016/S0022-5096(99)00028-9'),
    ('阅读范围', '题目、摘要、第 1 节（PDF 正文 pp.1–3；期刊 pp.797–799）'),
    ('本次目标', '搞清文章定位、能量泛函、演化条件、时间离散化和全文路线；不把第 2 节的公式推导记为已掌握'),
    ('一句话概括', '本文把 1998 年 Francfort–Marigo 的变分断裂模型转化为可计算的数值问题，用 Γ 收敛构造两种近似，并用三个算例检验模型的自洽性'),
    ('已理解', '裂纹是未知集合；总能量 = 表面能 + 弹性能；裂纹不可逆；每个载荷步在“包含上一步裂纹”的集合中最小化总能量；全局最小化可能给出最坏情况'),
    ('尚未理解', 'inf 与 min 的区别；条件 (c) 的选择作用；离散条件 (b) 的符号；全局/局部极小值对结果的影响；v 与 d 的符号约定'),
    ('下一步', '精读 §2.1 的 (1)–(3)，推导 E_c 的 Euler–Lagrange 方程和 v 的指数剖面；再读 §2.1.1 的交替最小化算法'),
], [1.35, 5.15], 9.0)

add_heading('二、逐段精读记录', 2)
add_table(['原文位置', '原文要点', '我的理解', '待确认'], [
    ('摘要 p.1', '给出 1998 模型的数值实现；提出基于变分近似的计算方法；在反平面和平面算例上测试；这些算例超出经典断裂力学计算工具的范围。', '本文的贡献是“算出来”，不是提出新模型；算例故意选择裂纹路径未知、可能分叉等经典方法难以处理的情形。', '“beyond the reach”具体指哪些情形？§3 用三个算例回答。'),
    ('§1 第 1 段 p.1', '跟随 1998 年的工作；不再重复讨论模型优缺点；目标是数值可行性研究；要让读者相信模型是可用、可管理的预测工具；选择“敌意”环境检验方法不会卡死。', '作者的目标是验证模型可计算、能给出符合理论框架的裂纹演化；这里的“predictive”是理论自洽意义上的，不是定量实验预测。', '“carefully tailored numerical version”需要哪些定制？§2 给出两条路线。'),
    ('§1 第 2 段 p.1', '仍达不到任何预测理论的终极目标，即实验结果的印证与预测；必须先做模型自洽性的内部审查；没有详细实验/数值协议时，对比只能是定性的。', '本文的评价标准是内部自洽 + 定性对比；读 §3 时不要把算例当成实验验证。作者明确说明这是模型自洽性研究。', '什么时候才能做定量实验对比？作者说需要详细的实验/数值协议。'),
    ('§1 第 3 段 pp.1–2', r'Ω 为均匀二维弹性体，A 为弹性张量，k 为断裂韧度；∂Ω\N 上给定位移 U(t,x)=tU(x)，N 上无面力；初始无裂纹；Γ 是 Ω̄ 的紧子集，可到达边界；E_s=kH¹(Γ\N)；E_e=inf_w ∫_{Ω\Γ} A e(w):e(w) dx；E=E_s+E_e。', '裂纹路径由表面能与弹性能的竞争决定；表面能是裂纹长度的线性函数；弹性能是在给定裂纹下放松位移场后的能量。', '为什么 E_e 是 inf 不是 min？Γ 不够光滑时 inf 可能取不到。'),
    ('演化条件 p.2', '(a) Γ(t) 随时间单调增，t<0 时为空；(b) 对任意包含全部历史裂纹的 Γ，E(Γ(t),t) ≤ E(Γ,t)；(c) 对任意 s<t，E(Γ(t),t) ≤ E(Γ(s),t)。', '(a) 不可逆；(b) 每步在“包含历史裂纹”的集合中全局最小；(c) 当前裂纹在当前载荷下不劣于任何过去状态，是选择准则。', '(c) 与 (b) 的区别？Γ(s) 不包含 s 之后的裂纹，所以不在 (b) 的容许集合中。'),
    ('全局/局部极小 p.2–3', '全局最小化可能过难且不现实；材料可能停在亚稳态；局部极小值可能更准确；能量非凸，数值格式可能检测到局部而非全局极小值。', '全局最小化是“最坏情况”；实际材料可能被能垒困住；数值结果可能依赖初值和网格。', '这会不会影响 §3 的裂纹路径和分叉结果？作者用非对称网格检查对称性。'),
    ('时间离散 p.3', 't0=0≤t1≤…；条件 (a)-(c) 退化为 (a) Γ_i ⊃ Γ_{i-1}；(b) E(Γ_i,t_i) ≤ E(Γ,t_i)，对任意 Γ ⊃ Γ_i（原文印刷）。', '每个载荷步在“包含上一步裂纹”的集合中最小化；但裂纹集合无法枚举，需要正则化。', '印刷的 Γ ⊃ Γ_i 是否应为 Γ ⊃ Γ_{i-1}？算法 p.9 是在 Γ ⊃ Γ_{i-1} 上最小化。'),
    ('全文路线 p.3', '方法 1：两场泛函 (u,v)，v 在裂纹附近为 0、远处为 1；方法 2：在三角剖分上定义单场泛函，第二场是网格；§2.3 比较；§2.4 平面弹性；§3 三个算例。', '两条路线都是“用可计算的场/网格隐式表示未知裂纹”；方法 1 是后来相场断裂 AT 正则化的祖先。', '方法 2 的“第二场是网格”如何理解？§2.2 详读。'),
    ('写作警告 pp.3–4', '第 2 节技术性强；只关心力学的读者可以跳过，但不要给第 2 节的近似赋予力学意义。', 'c、k_c 是数值参数，不是材料常数；但 c 可解释为损伤长度尺度，这是后来相场模型把 ℓ 当材料参数的源头之一。', '这与现代相场模型把 ℓ 当材料参数的做法有何区别？'),
], [0.95, 2.15, 2.15, 1.25], 8.4)

add_heading('三、符号与概念对照', 2)
add_table(['符号/概念', '含义', '易错点'], [
    ('Γ', '裂纹集合，Ω̄ 的紧子集，可到达边界', '不是预先给定的曲线；是待优化的未知集合。'),
    (r'H¹(Γ\N)', '二维中的裂纹长度；n−1 维 Hausdorff 测度', r'“fat crack”的 H¹ 为无穷，因此被 (b) 排除；减去 N 因为 traction-free 边界不增加表面能。'),
    ('k', '断裂韧度，单位裂纹长度的表面能', '与应力强度因子 K_IC 不同；这里是能量参数。'),
    ('A, e(w)', '弹性张量、小应变张量', 'e(w)=1/2(∇w+∇wᵀ)；A e:e 是弹性能密度（常数被吸收）。'),
    ('E_s, E_e, E', '表面能、弹性能、总能量', 'E_e 是 inf over w；E 是 Γ 的函数，时间 t 通过边界位移进入。'),
    ('U(t,x)=tU(x)', '比例加载', 't 是载荷因子；准静态，不是真实时间。'),
    ('Γ_i', '第 i 个载荷步的裂纹集合', '不可逆：Γ_i ⊃ Γ_{i-1}。'),
    ('v', '§2.1 的辅助场', '本文 v=1 完整、v=0 裂纹；现代相场常用 d=0 完整、d=1 裂纹，d=1−v。'),
    ('c, k_c, h', '正则化长度、残余刚度、网格尺寸', '理论上 h 比 k_c、c 更快趋于 0，且 k_c << c；实际算例 h=10^-3、c=10^-1、k_c=10^-4。裂纹附近需要 h << c。'),
    ('Γ 收敛', '泛函序列的变分收敛', '保证极小元的极限是极限泛函的极小元，但不给误差估计；本文用它证明正则化模型收敛到原模型。'),
], [1.15, 2.55, 2.80], 8.6)

add_heading('四、本次形成的五句话', 2)
for text in [
    '本文不是提出新模型，而是把 1998 年的变分断裂模型变成可计算的数值方法。',
    '总能量由表面能和弹性能组成；裂纹路径由两者竞争决定，不需要额外的裂纹扩展准则。',
    '裂纹不可逆；每个载荷步在“包含上一步裂纹”的集合中最小化总能量。',
    '裂纹集合无法直接枚举，因此用 Γ 收敛构造两种近似：连续场 v 的椭圆正则化，或自适应网格。',
    '全局最小化是“最坏情况”；实际材料可能停在亚稳态，数值结果也可能落在局部极小值。',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    style_run(r, 10.5)

add_heading('五、下次精读安排', 2)
add_body('下次从 §2.1 的公式 (1)–(3) 开始。重点做四件事：第一，弄清 F(u,v) 为什么要用 +∞ 扩展；第二，逐项解释 E_c 的两个积分；第三，由 E_c 对 v 变分，推导一维指数剖面 v=1−exp(−|x|/2c)，并核对表面能为什么是 1；第四，读懂 §2.1.1 的交替最小化算法和不可逆条件的实现方式。')

doc.save(path)
print('saved:', path)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables))
