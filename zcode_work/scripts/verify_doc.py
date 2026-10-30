from docx import Document
from pathlib import Path
p=Path(r'E:\909 研究\相场文献和程序\参考文献\相场文献阅读进展与问题记录.docx')
d=Document(p)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
for i,t in enumerate(d.tables): print('table',i,'rows',len(t.rows),'cols',len(t.columns))
print('title:',d.paragraphs[0].text)
print('last:',d.paragraphs[-1].text)
text='\n'.join(x.text for x in d.paragraphs)
print('contains_bad_replacement', '\ufffd' in text)
print('file_bytes',p.stat().st_size)
