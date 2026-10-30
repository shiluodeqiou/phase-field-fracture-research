from docx import Document
p=r'参考文献\相场文献阅读进展与问题记录.docx'; d=Document(p)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'shapes',len(d.inline_shapes))
print('last heading present', any(x.text=='第一篇阅读启动记录' for x in d.paragraphs))
print('last table rows',len(d.tables[-1].rows),'cols',len(d.tables[-1].columns))
