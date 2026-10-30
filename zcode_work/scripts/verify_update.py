from docx import Document
from pathlib import Path
p=Path(r'E:\909 研究\相场文献和程序\参考文献\相场文献阅读进展与问题记录.docx')
d=Document(p)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'inline_shapes',len(d.inline_shapes),'sections',len(d.sections))
print('table0 rows',len(d.tables[0].rows),'table5 rows',len(d.tables[5].rows))
for row in d.tables[0].rows[1:]: print(row.cells[0].text.replace('\n',' / '), '|', row.cells[3].text.replace('\n',' / '))
print('last_paragraph',d.paragraphs[-1].text)
print('size',p.stat().st_size)
